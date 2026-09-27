#Regression tests for pySupport
#
#The files in tests/fixtures only contain the sections of ORCA 6 and Gaussian 16 output files that
#pySupport reads. orca_tddft_triplet.out, orca_thermochemistry.out, gaussian_tddft_singlet.out,
#gaussian_tddft_triplet.out and gaussian_freq.out are trimmed real outputs (benzene, B3LYP/6-31G(d)),
#the other files are synthetic.
#
#Run from the repository root (Python >= 3.12):
#	python3 -m unittest discover -s tests -v

import contextlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)

import openpyxl
import pySupport
from modules import gaussian_analyzer, orca_analyzer, tex_generator, txt_generator, xlsx_generator
from modules.calc_data import CalcData

FIXTURES = os.path.join(REPO_ROOT, 'tests', 'fixtures')
ELEMENTS = ('H He Li Be B C N O F Ne Na Mg Al Si P S Cl Ar K Ca Sc Ti V Cr Mn Fe Co Ni Cu Zn Ga Ge As Se Br Kr '
	'Rb Sr Y Zr Nb Mo Tc Ru Rh Pd Ag Cd In Sn Sb Te I Xe Cs Ba La Ce Pr Nd Pm Sm Eu Gd Tb Dy Ho Er Tm Yb Lu '
	'Hf Ta W Re Os Ir Pt Au Hg Tl Pb Bi Po At Rn Fr Ra Ac Th Pa U Np Pu Am Cm Bk Cf Es Fm Md No Lr '
	'Rf Db Sg Bh Hs Mt Ds Rg Cn Nh Fl Mc Lv Ts Og').split()
WATER_ORCA = ['O      0.000000    0.000000    0.119262', 'H      0.000000    0.763239   -0.477047', 'H      0.000000   -0.763239   -0.477047']
WATER_GAUSSIAN = ['O 0.000000 0.000000 0.119262', 'H 0.000000 0.763239 -0.477047', 'H 0.000000 -0.763239 -0.477047']


def fixture(name):
	return os.path.join(FIXTURES, name)


def analyze(analyzer, path):
	#Run an analyzer without its console output
	with contextlib.redirect_stdout(io.StringIO()):
		return analyzer(path)


class TempDirTestCase(unittest.TestCase):
	#Tests that write files work in a temporary directory, pySupport writes its output next to the input files

	def setUp(self):
		tmp = tempfile.TemporaryDirectory()
		self.addCleanup(tmp.cleanup)
		self.addCleanup(os.chdir, os.getcwd())
		os.chdir(tmp.name)

	def write(self, name, text):
		with open(name, 'w') as f:
			f.write(text)
		return os.path.abspath(name)

	def copy_fixture(self, name, new_name=None):
		return shutil.copy(fixture(name), os.path.abspath(new_name or name))

	def modified_fixture(self, fixture_name, new_name, old, new):
		with open(fixture(fixture_name)) as f:
			text = f.read()
		self.assertIn(old, text)
		return self.write(new_name, text.replace(old, new))


class OrcaAnalyzerTest(unittest.TestCase):

	def test_opt_and_freq_in_one_keyword_line(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_opt_freq.out'))
		self.assertEqual(data.jobtype, 'opt+freq')
		self.assertEqual(data.imaginary_freqs, ['-150.00'])

	def test_optimization_uses_final_geometry(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_opt_freq.out'))
		self.assertEqual(data.coords, WATER_ORCA)

	def test_file_names_and_comments_do_not_change_job_type(self):
		#water_opt.xyz and "optimized" in a comment must not turn a frequency job into an optimization
		data = analyze(orca_analyzer.analyzer, fixture('orca_freq.out'))
		self.assertEqual(data.jobtype, 'freq')

	def test_frequency_job_has_coordinates(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_freq.out'))
		self.assertEqual(data.coords, WATER_ORCA)

	def test_tddft_single_point_has_coordinates(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_tddft.out'))
		self.assertEqual(data.jobtype, 'tddft')
		self.assertEqual(data.coords, WATER_ORCA)
		self.assertEqual(data.energies, [7.331939, 8.163398, 9.140216])

	def test_tddft_states_stay_aligned_with_their_energies(self):
		#State 2 has no contribution > 0.05 and keeps its largest one, the last line of state 3 must not be cut off
		data = analyze(orca_analyzer.analyzer, fixture('orca_tddft.out'))
		self.assertEqual(data.state_blocks, [[['4', '5', 0.99438]], [['2', '5', 0.04]], [['3', '5', 0.88], ['4', '6', 0.1]]])
		self.assertEqual(len(data.state_blocks), len(data.energies))

	def test_thermochemistry(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_thermochemistry.out'))
		self.assertEqual(data.jobtype, 'freq')
		self.assertEqual(data.thermochemistry, {
			'Sum of electronic and zero-point Energies': '-231.96771829',
			'Sum of electronic and thermal Energies': '-231.96324478',
			'Sum of electronic and thermal Enthalpies': '-231.96230057',
			'Sum of electronic and thermal Free Energies': '-231.99292080'})

	def test_homo_lumo_notation(self):
		#Benzene triplet: 22 alpha and 20 beta electrons, the notation refers to the HOMO of the same spin
		data = analyze(orca_analyzer.analyzer, fixture('orca_tddft_triplet.out'))
		self.assertEqual(data.homo, {'': 21, 'a': 21, 'b': 19})
		self.assertEqual([data.homo_lumo(contribution, '->') for contribution in data.state_blocks[0] + data.state_blocks[2] + data.state_blocks[3]], ['H -> L', 'H -> L', 'H-1 -> L', 'H -> L+1'])

	def test_tddft_open_shell_keeps_spin_labels(self):
		data = analyze(orca_analyzer.analyzer, fixture('orca_tddft_triplet.out'))
		self.assertEqual(data.state_blocks[0], [['21a', '22a', 0.493365], ['19b', '20b', 0.502866]])
		self.assertEqual(len(data.state_blocks), len(data.energies))


class GaussianAnalyzerTest(TempDirTestCase):

	def test_nosymm_uses_last_input_orientation(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_opt_nosymm.out'))
		self.assertEqual(data.coords, WATER_GAUSSIAN)

	def test_standard_orientation_is_preferred(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_genecp.out'))
		self.assertEqual(data.coords[0], 'Hf 0.000000 0.000000 0.000000')

	def test_general_basis_set(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_genecp.out'))
		self.assertEqual(data.basis_set, 'genecp')
		gen = self.modified_fixture('gaussian_genecp.out', 'gen.out', 'b3lyp/genecp', 'b3lyp/gen')
		self.assertEqual(analyze(gaussian_analyzer.analyzer, gen).basis_set, 'gen')

	def test_route_over_several_lines(self):
		#Long routes are wrapped after a fixed number of characters, also within a keyword (here Fr|eq)
		wrapped = self.modified_fixture('gaussian_freq.out', 'wrapped.out', ' #N B3LYP/6-31G(d) FREQ Geom=Connectivity\n',
			' #N B3LYP/6-31G(d) SCRF=(PCM,Solvent=Dichloromethane) Geom=Connectivity Fr\n eq\n')
		self.assertEqual(analyze(gaussian_analyzer.analyzer, wrapped).jobtype, 'freq')

	def test_job_type_from_keywords_only(self):
		#stable=opt is no geometry optimization and cphf=rdfreq no frequency calculation
		polar = self.modified_fixture('gaussian_freq.out', 'polar.out', ' #N B3LYP/6-31G(d) FREQ Geom=Connectivity\n', ' #N Stable=Opt B3LYP/6-31G(d) Polar CPHF=RdFreq\n')
		self.assertEqual(analyze(gaussian_analyzer.analyzer, polar).jobtype, 'other')

	def test_cis_is_an_excited_state_calculation(self):
		cis = self.modified_fixture('gaussian_tddft_singlet.out', 'cis.out', 'TD(NStates=20) B3LYP/6-31G(d)', 'CIS(NStates=20)/6-31G(d)')
		data = analyze(gaussian_analyzer.analyzer, cis)
		self.assertEqual(data.jobtype, 'tddft')
		self.assertEqual(len(data.state_blocks), 20)

	def test_frequencies_include_the_first_line(self):
		#The first "Frequencies --" line holds the lowest modes, i.e. the imaginary ones
		ts = self.modified_fixture('gaussian_freq.out', 'ts.out', 'Frequencies --    307.3016', 'Frequencies --   -307.3016')
		self.assertEqual(analyze(gaussian_analyzer.analyzer, ts).imaginary_freqs, ['-307.3016'])

	def test_thermochemistry(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_freq.out'))
		self.assertEqual(data.thermochemistry, {
			'Sum of electronic and zero-point Energies': '-232.127049',
			'Sum of electronic and thermal Energies': '-232.122582',
			'Sum of electronic and thermal Enthalpies': '-232.121637',
			'Sum of electronic and thermal Free Energies': '-232.152895'})

	def test_homo_lumo_notation(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_tddft_singlet.out'))
		self.assertEqual(data.homo[''], 20)
		self.assertEqual([data.homo_lumo(contribution, '->') for contribution in data.state_blocks[0]], ['H-1 -> L', 'H -> L+1'])

	def test_tddft_restricted_in_orca_format(self):
		#Orbitals counted from 0 and weights 2c², which add up to about 1 like in ORCA; weights <= 0.05 are filtered
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_tddft_singlet.out'))
		self.assertEqual(len(data.state_blocks), 20)
		self.assertEqual(len(data.energies), 20)
		blocks = [[[a, b, round(w, 4)] for a, b, w in block] for block in data.state_blocks]
		self.assertEqual(blocks[0], [['19', '21', 0.4994], ['20', '22', 0.4994]])
		self.assertEqual(blocks[4], [['20', '23', 0.995]])
		#18 -> 24 (0.12396) gives 2c² = 0.031 and is filtered
		self.assertEqual(blocks[14], [['16', '21', 0.8816], ['20', '28', 0.0813]])

	def test_tddft_unrestricted_in_orca_format(self):
		data = analyze(gaussian_analyzer.analyzer, fixture('gaussian_tddft_triplet.out'))
		self.assertEqual(len(data.state_blocks), 20)
		self.assertEqual(len(data.energies), 20)
		self.assertEqual(data.state_blocks[0], [['21a', '22a', 0.71274 ** 2], ['19b', '20b', 0.72053 ** 2]])

	def test_tddft_orbital_labels_match_orca(self):
		#Same molecule (benzene triplet, B3LYP/6-31G(d)): the lowest states have the same orbital labels in both programs
		labels = lambda data: [[entry[:2] for entry in block] for block in data.state_blocks[:5]]
		gaussian = analyze(gaussian_analyzer.analyzer, fixture('gaussian_tddft_triplet.out'))
		orca = analyze(orca_analyzer.analyzer, fixture('orca_tddft_triplet.out'))
		self.assertEqual(labels(gaussian), labels(orca))

	def test_element_symbols(self):
		rule = ' ' + '-' * 69
		lines = [
			' Entering Gaussian System, Link 0=g16',
			' #p b3lyp/6-31g(d)',
			' Charge =  0 Multiplicity = 1',
			'                         Standard orientation:',
			rule,
			' Center     Atomic      Atomic             Coordinates (Angstroms)',
			' Number     Number       Type             X           Y           Z',
			rule,
			*[f'{z:7d}{z:11d}{0:12d}{0:16.6f}{0:12.6f}{z:12.6f}' for z in range(1, 119)],
			rule,
			' Standard basis: 6-31G(d) (6D, 7F)',
			' SCF Done:  E(RB3LYP) =  -100.000000000     A.U. after   10 cycles',
			' Normal termination of Gaussian 16',
		]
		data = analyze(gaussian_analyzer.analyzer, self.write('elements.out', '\n'.join(lines) + '\n'))
		self.assertEqual([atom.split()[0] for atom in data.coords], ELEMENTS)


class PySupportTest(TempDirTestCase):

	def run_pysupport(self, si_style, *files):
		#Answer the style menu with si_style and every further question (TD-DFT summary) with the default, return the console output
		answers = iter([str(si_style)])
		with mock.patch.object(sys, 'argv', ['pySupport.py', *files]), mock.patch('builtins.input', side_effect=lambda prompt: next(answers, '')), contextlib.redirect_stdout(io.StringIO()) as output:
			pySupport.main()
		return output.getvalue()

	def test_unknown_and_missing_files_are_skipped(self):
		unknown = self.write('notes.out', 'not an ORCA or Gaussian output\n')
		output = self.run_pysupport(1, unknown, 'missing.out', self.copy_fixture('orca_opt_freq.out'))
		self.assertIn('Could not determine calculation program', output)
		self.assertTrue(os.path.isfile('orca_opt_freq.txt'))

	def test_abnormal_termination_skips_only_this_file(self):
		aborted = self.modified_fixture('orca_opt_freq.out', 'aborted.out', '****ORCA TERMINATED NORMALLY****', '')
		self.run_pysupport(1, aborted, self.copy_fixture('orca_freq.out'))
		self.assertFalse(os.path.exists('aborted.txt'))
		self.assertTrue(os.path.isfile('orca_freq.txt'))

	def test_excel_without_worksheets_is_not_written(self):
		self.run_pysupport(4, 'missing.out')
		self.assertFalse(os.path.exists('SI_output.xlsx'))

	def test_output_next_to_the_input_file(self):
		os.mkdir('calculations')
		freq = self.copy_fixture('orca_thermochemistry.out', os.path.join('calculations', 'freq.out'))
		self.run_pysupport(1, freq)
		self.assertTrue(os.path.isfile(os.path.join('calculations', 'freq.txt')))
		self.assertFalse(os.path.exists('freq.txt'))
		self.run_pysupport(4, freq)
		self.assertTrue(os.path.isfile(os.path.join('calculations', 'SI_output.xlsx')))
		self.assertFalse(os.path.exists('SI_output.xlsx'))

	def test_full_style_lists_the_thermochemistry(self):
		freq = self.copy_fixture('gaussian_freq.out')
		self.run_pysupport(1, freq)
		with open('gaussian_freq.txt') as f:
			self.assertIn('Sum of electronic and thermal Free Energies = -232.152895 Hartree\n', f.read())
		self.run_pysupport(4, freq)
		self.assertEqual(openpyxl.load_workbook('SI_output.xlsx')['gaussian_freq']['D9'].value, 'Sum of electronic and thermal Free Energies = -232.152895 Hartree')

	def test_excel_sheet_names_have_at_most_31_characters(self):
		#Both names start with the same 31 characters, the TD-DFT sheets also need room for "_TD-DFT"
		first = self.copy_fixture('orca_tddft.out', 'water_B3LYP_def2-SVP_TD-DFT_singlet_states_1.out')
		second = self.copy_fixture('orca_tddft.out', 'water_B3LYP_def2-SVP_TD-DFT_singlet_states_2.out')
		self.run_pysupport(4, first, second)
		sheetnames = openpyxl.load_workbook('SI_output.xlsx').sheetnames
		self.assertEqual(sheetnames, ['water_B3LYP_def2-SVP_TD-DFT_sin', 'water_B3LYP_def2-SVP_TD-_TD-DFT', 'water_B3LYP_def2-SVP_TD-DFT_s_2', 'water_B3LYP_def2-SVP_T_2_TD-DFT'])
		self.assertEqual(xlsx_generator.sheet_title(openpyxl.Workbook(), 'benzene[1]:opt'), 'benzene_1__opt')

	def test_excel_coordinates_only_skips_files_without_coordinates(self):
		no_coords = self.modified_fixture('gaussian_genecp.out', 'no_coords.out', 'orientation:', 'orientation removed')
		self.run_pysupport(6, no_coords, self.copy_fixture('orca_freq.out'))
		self.assertEqual(openpyxl.load_workbook('SI_output.xlsx').sheetnames, ['orca_freq'])


class TxtGeneratorTest(TempDirTestCase):

	def test_coordinate_columns_are_aligned(self):
		#Two-letter symbols and coordinates of 10 Angstroem and more must not shift the columns
		data = CalcData(os.path.abspath('aligned.out'), coords=['C -0.540860 0.467331 -0.297565', 'Cl 10.889091 -12.276545 2.493720'])
		with contextlib.redirect_stdout(io.StringIO()):
			txt_generator.generate_txt(3, data)
		with open('aligned.txt') as f:
			rows = [line for line in f.read().splitlines() if line.startswith('   C')]
		self.assertEqual(rows, ['   C      -0.540860       0.467331      -0.297565', '   Cl     10.889091     -12.276545       2.493720'])

	def test_tddft_summary_with_homo_lumo(self):
		data = analyze(orca_analyzer.analyzer, self.copy_fixture('orca_tddft_triplet.out'))
		with mock.patch('builtins.input', return_value=''), contextlib.redirect_stdout(io.StringIO()):
			txt_generator.generate_txt(2, data)
		with open('orca_tddft_triplet.txt') as f:
			text = f.read()
		self.assertIn('HOMO/LUMO', text)
		self.assertIn('3      18b -> 20b (0.994)      H-1 -> L      2.91         425.8            0.00\n', text)


class TexGeneratorTest(TempDirTestCase):

	def generate(self, si_style, name):
		#Write the .tex file for an ORCA fixture (TD-DFT summary: yes) and return its content
		data = analyze(orca_analyzer.analyzer, self.copy_fixture(name))
		with mock.patch('builtins.input', return_value='yes'), contextlib.redirect_stdout(io.StringIO()):
			tex_generator.generate_tex(si_style, data)
		with open(f'{name[:-3]}tex') as f:
			return f.read()

	def test_title_is_the_escaped_file_name(self):
		for si_style in (7, 8, 9):
			with self.subTest(si_style=si_style):
				tex = self.generate(si_style, 'orca_opt_freq.out')
				self.assertIn(r'\textbf{orca\_opt\_freq.out}', tex)
				self.assertNotIn('out.log', tex)

	def test_full_style_lists_the_thermochemistry(self):
		tex = self.generate(7, 'orca_thermochemistry.out')
		self.assertIn(r'\multicolumn{5}{l}{Sum of electronic and thermal Free Energies = -231.99292080 Hartree} \\', tex)

	def test_imaginary_frequencies_in_math_mode(self):
		tex = self.generate(7, 'orca_opt_freq.out')
		self.assertIn(r'$\nu_{i}$ = -150.00 cm$^{-1}$', tex)

	def test_tddft_summary_has_own_table(self):
		for si_style in (7, 8, 9):
			with self.subTest(si_style=si_style):
				coordinates, tddft = self.generate(si_style, 'orca_tddft.out').split(r'\begin{longtable}{cccccc}')
				self.assertIn('Cartesian Coordinates', coordinates)
				self.assertNotIn(r'\textbf{State}', coordinates)
				self.assertIn(r'\textbf{State}', tddft)
				self.assertIn(r'H-1 $\rightarrow$ L', tddft)

	@unittest.skipUnless(shutil.which('pdflatex'), 'pdflatex is not installed')
	def test_long_tables_break_across_pages(self):
		#150 atoms and 60 excited states do not fit on one page, nothing may be cut off at the bottom of a page
		coords = [f'C {n:.6f} 0.000000 0.000000' for n in range(150)]
		state_blocks = [[['10', '11', 0.9], ['9', '12', 0.08]] for n in range(60)]
		energies = [2.0 + n / 10 for n in range(60)]
		wavelengths = [1239.84 / energy for energy in energies]
		f_osc = ['0.10'] * 60
		for si_style in (7, 8, 9):
			with self.subTest(si_style=si_style):
				with mock.patch('builtins.input', return_value='yes'), contextlib.redirect_stdout(io.StringIO()):
					tex_generator.generate_tex(si_style, CalcData('large.out', 'def2-SVP', '0', '1', '-1000.0', 'tddft', coords=coords, homo={'': 10}, state_blocks=state_blocks, energies=energies, wavelengths=wavelengths, f_osc=f_osc))
				result = subprocess.run(['pdflatex', '-interaction=nonstopmode', '-halt-on-error', 'large.tex'], capture_output=True, text=True, errors='replace', timeout=120)
				self.assertEqual(result.returncode, 0, result.stdout[-1500:])
				with open('large.log', errors='replace') as f:
					self.assertNotIn(r'Overfull \vbox', f.read())

	@unittest.skipUnless(shutil.which('pdflatex'), 'pdflatex is not installed')
	def test_tex_files_compile(self):
		for name in ('orca_opt_freq.out', 'orca_tddft.out', 'orca_tddft_triplet.out', 'orca_thermochemistry.out'):
			for si_style in (7, 8, 9):
				with self.subTest(file=name, si_style=si_style):
					self.generate(si_style, name)
					result = subprocess.run(['pdflatex', '-interaction=nonstopmode', '-halt-on-error', f'{name[:-3]}tex'], capture_output=True, text=True, errors='replace', timeout=120)
					self.assertEqual(result.returncode, 0, result.stdout[-1500:])


if __name__ == '__main__':
	unittest.main()
